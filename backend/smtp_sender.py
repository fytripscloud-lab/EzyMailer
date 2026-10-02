"""Campaign sending over plain SMTP (stdlib `smtplib`), with no browser involved.

Each SMTP credential keeps one open connection for the whole campaign so a
lane doesn't pay a TLS handshake + AUTH per email; the caller owns that
connection object and closes it when the lane finishes.
"""
from __future__ import annotations

import csv
import re
import smtplib
import ssl
from email.utils import formatdate, make_msgid
from pathlib import Path

import certifi

from backend.gmail_oauth import _build_mime_message

DEFAULT_HOST = "smtp.gmail.com"
DEFAULT_PORT = 587
SECURITY_OPTIONS = ("Auto", "None", "SSL", "TLS", "TLS when available")
AUTO_RANDOM_SENDER_NAME = "Auto (Random Name)"
ACCOUNT_SENDER_NAME = "Google account name"
SENDER_NAME_OPTIONS = ("Auto", AUTO_RANDOM_SENDER_NAME, "$fullname", "$spanishname", "Custom")
# Gmail API accounts also have a real name on Google's side; keep it the default there.
API_SENDER_NAME_OPTIONS = (ACCOUNT_SENDER_NAME, *SENDER_NAME_OPTIONS)
SENDER_NAME_TOOLTIPS = {
    ACCOUNT_SENDER_NAME: "The Gmail account's own name from Google",
    "Auto": "A random full name picked once per account and kept for every campaign",
    AUTO_RANDOM_SENDER_NAME: "A new random full name for each account every campaign",
    "$fullname": "A random full name from the $fullname tag for each account every campaign",
    "$spanishname": "A random Spanish name from the $spanishname tag for each account every campaign",
    "Custom": "Type one sender name used by every account",
}

_SSL_CONTEXT = ssl.create_default_context(cafile=certifi.where())
_TIMEOUT_SECONDS = 30

# Google shows App Passwords as four space-separated groups of four letters.
_GOOGLE_APP_PASSWORD = re.compile(r"^[a-z]{4}( [a-z]{4}){3}$")
_CREDENTIAL_LINE = re.compile(r"^\s*([^\s,;:|]+@[^\s,;:|]+)\s*[,;:|\t ]\s*(.+?)\s*$")


class SmtpAuthError(RuntimeError):
    """The server rejected the username/password; retrying won't help."""


def resolve_security(security: str, port: int) -> str:
    """Map the UI choice to one concrete transport: ssl, starttls, starttls-optional or none.

    Auto follows the port's convention: 465 is implicit TLS, 587 is the
    submission port where STARTTLS is mandatory, anything else upgrades to
    TLS only if the server offers it.
    """
    if security == "SSL":
        return "ssl"
    if security == "TLS":
        return "starttls"
    if security == "TLS when available":
        return "starttls-optional"
    if security == "None":
        return "none"
    if port == 465:
        return "ssl"
    if port == 587:
        return "starttls"
    return "starttls-optional"


def connect(host: str, port: int, security: str, username: str, password: str) -> smtplib.SMTP:
    transport = resolve_security(security, port)
    try:
        if transport == "ssl":
            client: smtplib.SMTP = smtplib.SMTP_SSL(host, port, timeout=_TIMEOUT_SECONDS, context=_SSL_CONTEXT)
        else:
            client = smtplib.SMTP(host, port, timeout=_TIMEOUT_SECONDS)
        client.ehlo()
        if transport in {"starttls", "starttls-optional"}:
            if client.has_extn("starttls"):
                client.starttls(context=_SSL_CONTEXT)
                client.ehlo()
            elif transport == "starttls":
                client.close()
                raise RuntimeError(f"{host}:{port} does not support TLS. Choose another Security option.")
    except (OSError, smtplib.SMTPException) as exc:
        raise RuntimeError(f"Could not connect to {host}:{port} — {exc}") from exc

    if password:
        try:
            client.login(username, password)
        except smtplib.SMTPAuthenticationError as exc:
            close(client)
            detail = exc.smtp_error.decode("utf-8", errors="replace") if isinstance(exc.smtp_error, bytes) else str(exc.smtp_error)
            hint = " Gmail needs an App Password, not the normal account password." if "gmail" in host.lower() else ""
            raise SmtpAuthError(f"Login rejected for {username} ({exc.smtp_code}): {detail.strip()}.{hint}") from exc
        except smtplib.SMTPNotSupportedError as exc:
            close(client)
            raise RuntimeError(
                f"{host}:{port} does not allow login without TLS. Set Security to Auto, SSL or TLS."
            ) from exc
    return client


def close(client: smtplib.SMTP | None) -> None:
    if client is None:
        return
    try:
        client.quit()
    except Exception:
        try:
            client.close()
        except Exception:
            pass


def send_message(
    client: smtplib.SMTP,
    sender_email: str,
    to: str,
    subject: str,
    body_text: str,
    *,
    sender_name: str = "",
    html_body: str | None = None,
    attachment_paths: list[Path] | None = None,
) -> None:
    domain = sender_email.rsplit("@", 1)[-1] if "@" in sender_email else None
    raw = _build_mime_message(
        sender_email,
        sender_name,
        to,
        subject,
        body_text,
        html_body,
        attachment_paths or [],
        extra_headers={"Date": formatdate(localtime=True), "Message-ID": make_msgid(domain=domain)},
    )
    refused = client.sendmail(sender_email, [to], raw)
    if refused:
        code, reason = refused.get(to, (0, b""))
        reason_text = reason.decode("utf-8", errors="replace") if isinstance(reason, bytes) else str(reason)
        raise RuntimeError(f"Server refused {to} ({code}): {reason_text}")


def is_connection_error(exc: BaseException) -> bool:
    """True when the open connection is unusable and a fresh one may succeed."""
    if isinstance(exc, smtplib.SMTPServerDisconnected):
        return True
    # smtplib.SMTPException subclasses OSError; those are protocol replies, not dropped sockets.
    return isinstance(exc, OSError) and not isinstance(exc, smtplib.SMTPException)


CREDENTIAL_FILE_SUFFIXES = (".csv", ".xlsx", ".xlsm", ".xls", ".txt")
SAMPLE_CREDENTIALS_CSV = (
    "email,password\r\n"
    "sender1@gmail.com,abcdefghijklmnop\r\n"
    "sender2@yourdomain.com,your-smtp-password\r\n"
)


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    # Excel stores an all-digit password as a number; 123456 must not become "123456.0".
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _credentials_from_rows(rows) -> list[tuple[str, str]]:
    """First cell containing "@" is the email, the next cell its password.

    A header row has no "@" so it is skipped; a single "email:password"
    cell is split the same way text files are.
    """
    credentials: list[tuple[str, str]] = []
    for raw_row in rows:
        cells = [_cell_text(value) for value in raw_row]
        for index, cell in enumerate(cells):
            if "@" not in cell:
                continue
            password = cells[index + 1] if index + 1 < len(cells) else ""
            if password:
                if _GOOGLE_APP_PASSWORD.match(password):
                    password = password.replace(" ", "")
                credentials.append((cell, password))
            else:
                credentials.extend(parse_credentials(cell))
            break
    return credentials


def read_credentials_file(path: Path) -> list[tuple[str, str]]:
    """Read (email, password) pairs from a .csv, .xlsx/.xlsm, .xls or .txt file (first sheet only)."""
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            return _credentials_from_rows(workbook.active.iter_rows(values_only=True))
        finally:
            workbook.close()
    if suffix == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(path))
        sheet = book.sheet_by_index(0)
        return _credentials_from_rows(sheet.row_values(index) for index in range(sheet.nrows))
    return parse_credentials(path.read_text(encoding="utf-8-sig", errors="replace"))


def parse_credentials(text: str) -> list[tuple[str, str]]:
    """Read `email<sep>password` lines (sep: comma, colon, semicolon, pipe, tab or space).

    Lines without an email address — headers, blanks, comments — are skipped.
    The password is everything after the first separator, so passwords that
    themselves contain separators survive intact.
    """
    credentials: list[tuple[str, str]] = []
    for raw_line in text.splitlines():
        line = raw_line.strip().lstrip("﻿")
        if not line or line.startswith("#"):
            continue
        email = password = ""
        if '"' in line:
            fields = [field.strip() for field in next(csv.reader([line]), [])]
            if len(fields) >= 2 and "@" in fields[0]:
                email, password = fields[0], fields[1]
        if not email:
            match = _CREDENTIAL_LINE.match(line)
            if not match:
                continue
            email, password = match.group(1), match.group(2)
        if _GOOGLE_APP_PASSWORD.match(password):
            password = password.replace(" ", "")
        if email and password:
            credentials.append((email.strip(), password))
    return credentials
